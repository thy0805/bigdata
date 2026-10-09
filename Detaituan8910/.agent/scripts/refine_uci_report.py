from pathlib import Path
import json
from docx import Document
from docx.shared import Cm
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

project = Path('D:/Hoctap/bigdata/Detaituan8910')
path = project / 'BaoCao_PhanTich_DuDoan_DienNang_UCI_20261008.docx'
doc = Document(path)
fields = 0
for node in doc.element.xpath('.//w:instrText'):
    if 'STYLEREF' in (node.text or ''):
        node.text = node.text.replace('\\n', '\\s')
        fields += 1
for table, widths in zip(doc.tables[:2], ([1.3, 4.4, 3.05, 3.5, 2.75], [1.15, 2.5, 3.55, 5.3, 2.5])):
    for col, width in zip(table.columns, widths):
        col.width = Cm(width)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Cm(width)
for p in doc.paragraphs:
    if p.style.style_id == 'TableCaption':
        p.paragraph_format.keep_with_next = True
    if p.style.style_id == 'Heading21' and p.text.endswith('(kWh)'):
        value = p.text
        p.clear()
        p.add_run(value[:-5])
        p.add_run('(kWh)').font.all_caps = False
    if p.style.style_id in ('TOC1', 'TOC2', 'TOC3'):
        p.paragraph_format.right_indent = Cm(1.4)
    if p.text.startswith('[') and p.style.style_id == 'ReportBody':
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
doc.save(path)
print(json.dumps({'caption_fields_changed': fields}, ensure_ascii=False))
