from pathlib import Path
import hashlib
import json
import re
import zipfile
from docx import Document
from docx.oxml.ns import qn, nsmap

project = Path('D:/Hoctap/bigdata/Detaituan8910')
qa = project / '.agent/qa/word-uci-20261008'
output = project / 'BaoCao_PhanTich_DuDoan_DienNang_UCI_20261008.docx'
manifest = json.loads((qa / 'build-manifest.json').read_text(encoding='utf-8-sig'))
word = json.loads((qa / 'final-word.json').read_text(encoding='utf-8-sig'))
probe = json.loads((qa / 'final-field-tests.json').read_text(encoding='utf-8-sig'))
dataset = json.loads((qa / 'dataset-inspection.json').read_text(encoding='utf-8-sig'))
pdf = json.loads((qa / 'final-pdf.json').read_text(encoding='utf-8-sig'))
doc = Document(output)
source = Document(project / 'KhoiTuan Tong Quan Bai Toan Chuong 1.docx')
checks = {}
details = {}


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def norm(value):
    return re.sub(r'\s+', ' ', re.sub(r'\[\d+\]', '', value)).replace(' .', '.').strip()


def xml_text(node):
    return ''.join('\n' if el.tag == qn('w:br') else (el.text or '') for el in node.iter() if el.tag in (qn('w:t'), qn('w:br')))


hashes = {path: sha(path) for path in manifest['source_hashes']}
checks['four_docx_sources_unchanged'] = hashes == manifest['source_hashes']
checks['teacher_pdf_unchanged'] = sha(project / 'Ke hoach Do an mon hoc Nhap mon Big data.pdf') == '61e91b9c3822994a1822de9f7f51b58f97842f3d5b1f9dfcf792f23c0feeeea4'
checks['dataset_zip_unchanged'] = sha(dataset['path']) == dataset['sha256']
checks['dataset_structure'] = dataset['rows'] == 2075259 and len(dataset['columns']) == 9 and dataset['row_width_counts'] == {'9': 2075259}
checks['pages_and_render'] = word['pages'] == len(pdf) == 37 and all((qa / f'final-render/page-{n}.png').exists() for n in range(1, 38))
checks['heading_count'] = len(word['headings']) == 92
chapter_headings = [h for h in word['headings'] if h['style'] == 'ChapterTitle']
checks['five_chapter_numbers'] = [h['number'] for h in chapter_headings] == [f'CHƯƠNG {n}.' for n in range(1, 6)]
h2 = [h for h in word['headings'] if h['style'] == 'Heading 21']
expected_numbers = [f'{chapter}.{index}.' for chapter, count in enumerate([11, 14, 13, 14, 17], 1) for index in range(1, count + 1)]
checks['all_69_chapter_sections_numbered'] = [h['number'] for h in h2] == expected_numbers
checks['77_second_level_headings'] = len([h for h in word['headings'] if h['level'] == 2]) == 77
chapter2_source = Document(project / 'HAULYVUNHAN_Chuong2_CoSoLyThuyet.docx')
outline2 = [re.sub(r'^2\.\d+\.\s*', '', p.text).strip().upper() for p in chapter2_source.paragraphs if re.match(r'^2\.\d+\.', p.text)]
checks['chapter2_exact_outline'] = outline2 == [h['text'].upper() for h in h2 if h['number'].startswith('2.')]
chapter_index = 0
nonheading_after_ch1 = []
chapter1_paragraphs = []
for p in doc.paragraphs:
    if p.style.style_id == 'ChapterTitle':
        chapter_index += 1
    if p.text == 'KẾT LUẬN':
        chapter_index = 6
    if chapter_index == 1 and p.style.style_id == 'ReportBody':
        chapter1_paragraphs.append(p.text)
    if 2 <= chapter_index <= 5 and p.text.strip() and p.style.style_id not in ('ChapterTitle', 'Heading21'):
        nonheading_after_ch1.append((chapter_index, p.style.style_id, p.text))
checks['chapters2_to5_no_invented_prose'] = not nonheading_after_ch1
checks['uci_and_hour_ahead_outline'] = all(any(term in h['text'].lower() for h in h2) for term in ['uci individual household', 'giờ tiếp theo', '(kwh)', 'cửa sổ thời gian một giờ', 'mô hình đã huấn luyện'])
final_ch1_normalized = [norm(t) for t in chapter1_paragraphs]
source_body_checked = []
exceptions = {13: 'Chú thích nguồn bảng', 54: 'Thu hẹp kết luận Gasparin theo paper', 66: 'Chốt Flink và tách huấn luyện mô hình'}
retained = []
for item in manifest['changes']:
    index = item['source_block']
    expected = item['text']
    source_lines = [line.strip() for line in xml_text(source.element.body[index]).split('\n') if line.strip() and not re.match(r'^1\.\d+\.', line.strip())]
    source_matches = any(norm(line).rstrip('.') == norm(expected).rstrip('.') for line in source_lines)
    final_matches = norm(expected) in final_ch1_normalized
    source_body_checked.append({'source_block': index, 'present_in_final': final_matches, 'same_source_text_except_citations': source_matches, 'allowed_change': exceptions.get(index)})
    retained.append(final_matches and (source_matches or index in exceptions))
checks['all_55_source_body_paragraphs_preserved_or_targeted'] = len(retained) == 55 and all(retained)
checks['added_1_3'] = any(h['number'] == '1.3.' and h['text'] == 'DỮ LIỆU CÔNG TƠ THÔNG MINH' for h in h2) and any('Advanced Metering Infrastructure' in t for t in chapter1_paragraphs)
checks['kW_kWh_distinction_and_missing_hour_limit'] = any('Không được đổi tên trực tiếp cột công suất kW thành điện năng kWh' in t for t in chapter1_paragraphs) and any('Chỉ những giờ có đủ bản ghi hợp lệ' in t for t in chapter1_paragraphs)
checks['seven_tables_and_four_ch1_tables'] = len(doc.tables) == 7 and len(source.tables) == 4
table_comparisons = []
for number, (old, new) in enumerate(zip(source.tables, doc.tables[3:]), 1):
    matches = []
    for r, row in enumerate(old.rows):
        for c, cell in enumerate(row.cells):
            changed_example = number == 4 and c == 2 and r in (1, 2)
            matches.append(norm(cell.text) == norm(new.cell(r, c).text) or changed_example)
    table_comparisons.append(all(matches))
checks['source_table_cells_retained_except_two_tech_examples'] = all(table_comparisons)
checks['all_tables_repeat_header'] = all(t.rows[0]._tr.xpath('./w:trPr/w:tblHeader') for t in doc.tables)
with zipfile.ZipFile(output) as package:
    media_hashes = {hashlib.sha256(package.read(n)).hexdigest() for n in package.namelist() if n.startswith('word/media/')}
checks['source_figure_bytes_preserved'] = set(manifest['source_image_sha256']).issubset(media_hashes)
template = Document('D:/Hoctap/bigdata/Mauwword/Mau bao cao_Do an_Khoa luan_2025.docx')
template_images = [hashlib.sha256(r.target_part.blob).hexdigest() for r in template.part.rels.values() if r.reltype.endswith('/image')]
checks['template_logo_present'] = bool(set(template_images) & media_hashes)
checks['three_inline_images'] = len(doc.inline_shapes) == 3
maths = doc.element.xpath('.//m:oMath')
checks['three_native_equations'] = len(maths) == 3
checks['equation_structure'] = bool(maths[0].xpath('./m:sSub/m:e/m:acc', namespaces=nsmap)) and bool(maths[1].xpath('./m:f', namespaces=nsmap)) and bool(maths[1].xpath('./m:nary', namespaces=nsmap)) and bool(maths[2].xpath('./m:rad/m:e/m:nary/m:e/m:sSup', namespaces=nsmap))
field_codes = [el.text or '' for el in doc.element.xpath('.//w:instrText')] + [el.get(qn('w:instr'), '') for el in doc.element.xpath('.//w:fldSimple')]
checks['three_real_toc_fields'] = len([s for s in field_codes if s.strip().startswith('TOC ')]) == 3
checks['ten_caption_fields'] = len([s for s in field_codes if 'STYLEREF' in s or s.strip().startswith('SEQ ')]) == 10
checks['main_toc_has_91_entries'] = len([x for x in word['tocs'][0].split('\r') if x.strip()]) == 91
checks['caption_lists_correct'] = len([x for x in word['tocs'][1].split('\r') if x.strip()]) == 1 and len([x for x in word['tocs'][2].split('\r') if x.strip()]) == 4 and 'Hình 1.1.' in word['tocs'][1] and all(f'Bảng 1.{n}.' in word['tocs'][2] for n in range(1, 5))
checks['native_word_insert_delete_probe'] = all(probe[k] for k in ('toc_detects_inserted', 'toc_detects_h3', 'figure_list_detects_caption', 'table_list_detects_caption')) and probe['changes_saved'] is False and [probe[k] for k in ('inserted_h2', 'following_h2', 'restored_h2', 'inserted_h3')] == ['3.2.', '3.3.', '3.2.', '3.1.1.']
links = {r.target_ref for r in doc.part.rels.values() if r.reltype.endswith('/hyperlink') and r.is_external}
checks['16_clickable_reference_links'] = all(ref[-1] in links for ref in manifest['refs']) and len(manifest['refs']) == 16
reference_numbers = [int(re.match(r'^\[(\d+)\]', p.text).group(1)) for p in doc.paragraphs if re.match(r'^\[\d+\]', p.text)]
checks['references_numbered1_to16'] = reference_numbers == list(range(1, 17))
text_all = '\n'.join(xml_text(n) for n in doc.element.body)
checks['no_london_or_thy_or_tool_tokens'] = all(s not in text_all for s in ['London', 'Quách Bảo Thy', 'chatgpt-content-reference', 'turn74view', 'Kiểm thử chèn', 'Kiểm thử danh mục'])
checks['three_confirmed_mssv'] = all(s in text_all for s in ['2001230640', '2001230227', '2001230418'])
checks['three_students_and_uninvented_contribution'] = len(doc.tables[1].rows) == 4 and all(not row.cells[-1].text.strip() for row in doc.tables[1].rows[1:])
checks['blank_weekly_history'] = len(doc.tables[0].rows) == 11 and all(not cell.text.strip() for row in doc.tables[0].rows[1:] for cell in row.cells[1:])
checks['a4_margins'] = all(abs(s.page_width.cm - 21) < .02 and abs(s.page_height.cm - 29.7) < .02 and abs(s.left_margin.cm - 3.5) < .02 and abs(s.right_margin.cm - 2.5) < .02 and abs(s.top_margin.cm - 2.5) < .02 and abs(s.bottom_margin.cm - 2.5) < .02 for s in doc.sections)
default_font = doc.styles.element.xpath('./w:docDefaults/w:rPrDefault/w:rPr/w:rFonts')[0].get(qn('w:ascii'))
default_size = int(doc.styles.element.xpath('./w:docDefaults/w:rPrDefault/w:rPr/w:sz')[0].get(qn('w:val'))) / 2
body_font = doc.styles['ReportBody'].font.name or doc.styles['Normal'].font.name or default_font
body_size = doc.styles['ReportBody'].font.size or doc.styles['Normal'].font.size
body_size_pt = body_size.pt if body_size else default_size
body_spacing = doc.styles['ReportBody'].paragraph_format.line_spacing or doc.styles['Normal'].paragraph_format.line_spacing
checks['report_body_format'] = body_font == 'Times New Roman' and body_size_pt == 13 and abs(body_spacing - 1.3) < .01
number_formats = [s._sectPr.xpath('./w:pgNumType')[0].get(qn('w:fmt')) if s._sectPr.xpath('./w:pgNumType') else None for s in doc.sections]
checks['front_roman_body_arabic'] = number_formats[1] == 'lowerRoman' and number_formats[2] in (None, 'decimal') and all(s._sectPr.xpath('./w:pgNumType')[0].get(qn('w:start')) == '1' for s in list(doc.sections)[1:])
checks['covers_no_page_number'] = all(not h.strip().endswith('PAGE') for h in [el.text or '' for el in doc.sections[0].footer._element.xpath('.//w:instrText')]) and pdf[0]['text'].strip().splitlines()[-1].strip() not in ['1', 'i'] and pdf[1]['text'].strip().splitlines()[-1].strip() not in ['2', 'ii']
checks['first_cover_frame'] = bool(doc.sections[0]._sectPr.xpath('./w:pgBorders/w:top'))
checks['body_starts_page1'] = any(h['text'] == 'MỞ ĐẦU' and h['page'] == 1 and h['physical_page'] == 13 for h in word['headings'])
checks['no_field_errors'] = all(not re.search(r'Error!|Lỗi!|Bookmark not defined|Reference source not found', page['text'], re.I) for page in pdf)
details.update({'source_hashes': hashes, 'zip_sha256': dataset['sha256'], 'source_paragraph_comparisons': source_body_checked, 'source_table_comparisons': table_comparisons, 'source_figure_sha256': manifest['source_image_sha256'], 'pages': 37, 'headings': 92, 'toc_entries': 91, 'caption_lists': word['tocs'][1:], 'page_number_formats': number_formats, 'pending': ['Chương 2 chưa có nội dung ngoài đề cương', 'Chương 3–5 chưa thực nghiệm/triển khai', 'Lời cảm ơn, Mở đầu, Kết luận vẫn là khung', 'Lịch tuần, phần việc người thứ ba và mức đóng góp cần thông tin thật'], 'reference_network_limitations': '13 links HTTP 200; IEEE challenge 202, Wiley automation 403, JKU TLS issue in Python. Primary JKU PDF and Wiley article were retrieved through web tools; no dead-link conclusion from access limitations.'})
result = {'status': 'STRUCTURAL_SEMANTIC_RUNTIME_PASS' if all(checks.values()) else 'FAILED', 'output': str(output), 'output_sha256': sha(output), 'checks': checks, 'details': details}
(qa / 'final-verification.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps({'status': result['status'], 'checks_passed': sum(checks.values()), 'checks_total': len(checks), 'failed': [k for k, v in checks.items() if not v], 'sha256': result['output_sha256']}, ensure_ascii=False))
if not all(checks.values()):
    raise SystemExit(1)
