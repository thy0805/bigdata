from copy import deepcopy
from io import BytesIO
from pathlib import Path
import hashlib
import json
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Cm, Pt, RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph


ROOT = Path('D:/Hoctap/bigdata')
PROJECT = ROOT / 'Detaituan8910'
QA = PROJECT / '.agent/qa/word-uci-20261008'
SKELETON = ROOT / 'BaoCao_PhanTich_DuDoan_DienNang_Khung_v2.docx'
TEMPLATE = ROOT / 'Mauwword/Mau bao cao_Do an_Khoa luan_2025.docx'
CHAPTER = PROJECT / 'KhoiTuan Tong Quan Bai Toan Chuong 1.docx'
CHAPTER2 = PROJECT / 'HAULYVUNHAN_Chuong2_CoSoLyThuyet.docx'
OUTPUT = Path(sys.argv[1]) if len(sys.argv) > 1 else PROJECT / 'BaoCao_PhanTich_DuDoan_DienNang_UCI_20261008.docx'
TITLE = 'PHÂN TÍCH DỮ LIỆU TIÊU THỤ ĐIỆN NĂNG\nVÀ DỰ ĐOÁN NHU CẦU SỬ DỤNG ĐIỆN THEO THỜI GIAN'
MEMBERS = [('2001230640', 'Nguyễn Đức Thành Phát'), ('2001230227', 'Lý Vũ Nhân Hậu'), ('2001230418', 'Nguyễn Tuấn Khôi')]


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def text(el):
    return ''.join((n.text or '') if n.tag == qn('w:t') else '\n' if n.tag == qn('w:br') else '' for n in el.iter())


def put(el, tag, **attrs):
    node = el.find(qn(tag))
    if node is None:
        node = OxmlElement(tag)
        el.append(node)
    for key, value in attrs.items():
        node.set(qn(key), str(value))
    return node


if OUTPUT.exists():
    raise FileExistsError(OUTPUT)
source_hashes = {str(p): sha(p) for p in (SKELETON, TEMPLATE, CHAPTER, CHAPTER2)}
doc = Document(SKELETON)
template = Document(TEMPLATE)
chapter = Document(CHAPTER)
original = list(doc.element.body)
cover_section = deepcopy(original[14].find(qn('w:pPr')).find(qn('w:sectPr')))
front_section = deepcopy(original[121].find(qn('w:pPr')).find(qn('w:sectPr')))
body_section = deepcopy(original[-1])
chapter_headers = [deepcopy(e) for e in original if e.tag == qn('w:p') and e.xpath('./w:pPr/w:pStyle[@w:val="ChapterTitle"]')]
h2_prototype = deepcopy(original[134])
intro_prototypes = [deepcopy(original[i]) for i in (123, 125, 127, 129, 131)]
conclusion_prototypes = [deepcopy(original[i]) for i in (277, 279, 281)]
chapter1_titles = [text(original[i]) for i in range(134, 156, 2)]
tail = [deepcopy(e) for e in original[156:285]]
for e in list(doc.element.body):
    doc.element.body.remove(e)
doc.element.body.append(body_section)


def set_font(run, size=None, bold=None, italic=None):
    run.font.name = 'Times New Roman'
    run.font.color.rgb = RGBColor(0, 0, 0)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    rf = run._r.get_or_add_rPr().get_or_add_rFonts()
    for k in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        rf.set(qn('w:' + k), 'Times New Roman')
    for k in list(rf.attrib):
        if 'Theme' in k:
            del rf.attrib[k]
    color = run._r.get_or_add_rPr().find(qn('w:color'))
    if color is not None:
        for k in list(color.attrib):
            if k != qn('w:val'):
                del color.attrib[k]


def p(text_value='', style='ReportBody'):
    para = doc.add_paragraph(style=style)
    if text_value:
        set_font(para.add_run(text_value))
    return para


def append(el):
    doc.element.body.insert(len(doc.element.body) - 1, el)
    return el


def remap_images(el, source):
    for blip in el.xpath('.//a:blip'):
        old = blip.get(qn('r:embed'))
        if old:
            blob = source.part.rels[old].target_part.blob
            rid, _ = doc.part.get_or_add_image(BytesIO(blob))
            blip.set(qn('r:embed'), rid)
    for node in el.xpath('.//wp:docPr'):
        node.set('id', str(100 + len(doc.element.xpath('.//wp:docPr'))))


def field(para, instruction, cached=''):
    for kind in ('begin', 'separate', 'end'):
        r = OxmlElement('w:r')
        n = OxmlElement('w:fldChar')
        n.set(qn('w:fldCharType'), kind)
        if kind == 'begin':
            n.set(qn('w:dirty'), 'true')
        r.append(n)
        para._p.append(r)
        if kind == 'begin':
            r = OxmlElement('w:r')
            n = OxmlElement('w:instrText')
            n.set(qn('xml:space'), 'preserve')
            n.text = ' ' + instruction + ' '
            r.append(n)
            para._p.append(r)
        if kind == 'separate' and cached:
            set_font(para.add_run(cached))


def section_break(sect):
    para = p()
    pf = para.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = 1
    pf.keep_with_next = False
    para._p.get_or_add_pPr().append(deepcopy(sect))
    put(para._p.get_or_add_pPr(), 'w:rPr').append(OxmlElement('w:sz'))
    para._p.find(qn('w:pPr')).find(qn('w:rPr')).find(qn('w:sz')).set(qn('w:val'), '2')


for style in doc.styles:
    if style.type != 1:
        continue
    style.font.name = 'Times New Roman'
    style.font.color.rgb = RGBColor(0, 0, 0)
    rf = style.element.get_or_add_rPr().get_or_add_rFonts()
    for key in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        rf.set(qn('w:' + key), 'Times New Roman')
    for node in list(style.element.xpath('./w:pPr/w:pBdr')):
        node.getparent().remove(node)

for name in ('Normal', 'ReportBody'):
    s = doc.styles[name]
    s.font.size = Pt(13)
    f = s.paragraph_format
    f.line_spacing = 1.3
    f.space_before = Pt(6)
    f.space_after = Pt(6)
    f.first_line_indent = Cm(.75)
    f.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    f.widow_control = True
    f.keep_with_next = False
for name, size in (('ChapterTitle', 18), ('Heading 11', 18), ('Heading 21', 14), ('Heading 31', 13), ('FrontTitle', 18), ('FrontTitleNoTOC', 18)):
    s = doc.styles[name]
    s.font.size = Pt(size)
    s.font.bold = True
    s.font.all_caps = name == 'Heading 21'
    s.paragraph_format.first_line_indent = Cm(0)
    s.paragraph_format.keep_with_next = True
    s.paragraph_format.keep_together = True
for name in ('FigureCaption', 'TableCaption'):
    s = doc.styles[name]
    s.font.size = Pt(12)
    s.paragraph_format.first_line_indent = Cm(0)
    s.paragraph_format.keep_together = True
for name in ('toc 1', 'toc 2', 'toc 3'):
    if name in doc.styles:
        s = doc.styles[name]
        s.font.size = Pt(13)
        s.paragraph_format.line_spacing = 1.3
        s.paragraph_format.space_before = Pt(0)
        s.paragraph_format.space_after = Pt(0)
        s.paragraph_format.first_line_indent = Cm(0)

cover_indices = ([0, 1, 2, 3, 4, 7, 10, 12, 15, 16, 17, 18, 23], [24, 25, 26, 27, 28, 31, 34, 36, 38, 39, 41, 42, 43, 44, 47])
for cover_no, indices in enumerate(cover_indices):
    for index in indices:
        el = deepcopy(template.element.body[index])
        for node in el.xpath('./w:pPr/w:sectPr'):
            node.getparent().remove(node)
        remap_images(el, template)
        para = Paragraph(append(el), doc._body)
        old = para.text
        if index in (7, 31):
            para.text = 'BÀI TẬP NHÓM MÔN BIGDATA'
        elif index in (10, 34):
            para.text = TITLE
        elif index in (12, 36):
            para.text = 'Ngành: Công nghệ thông tin'
        elif index == 39:
            para.text = 'Nguyễn Thành Ngô'
        elif index in (16, 17, 18, 42, 43, 44):
            offset = index - (16 if cover_no == 0 else 42)
            mssv, name = MEMBERS[offset]
            para.text = f'{offset + 1}. {mssv} – {name}'
        elif index in (23, 47):
            para.text = 'TP. HỒ CHÍ MINH, tháng 10 năm 2026'
        pf = para.paragraph_format
        pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf.first_line_indent = Cm(0)
        pf.left_indent = Cm(0)
        pf.right_indent = Cm(0)
        pf.space_before = Pt(0)
        pf.space_after = Pt(6)
        pf.line_spacing = 1.15
        pf.keep_with_next = False
        pf.keep_together = True
        pf.page_break_before = index == 24
        size = 14
        bold = index in (1, 2, 25, 26, 7, 31, 10, 34, 15, 41, 38)
        if index in (7, 31):
            size = 16
            pf.space_before = Pt(24 if cover_no == 0 else 12)
            pf.space_after = Pt(24 if cover_no == 0 else 12)
        if index in (10, 34):
            size = 18
            pf.line_spacing = 1.3
            pf.space_after = Pt(18)
        if index in (15, 38):
            pf.space_before = Pt(24 if cover_no == 0 else 12)
        if index == 41:
            pf.space_before = Pt(12)
        if index in (23, 47):
            pf.space_before = Pt(55 if cover_no == 0 else 30)
        for run in para.runs:
            if not run._r.xpath('.//w:sym'):
                set_font(run, size, bold, False)
        if old == '------':
            pf.space_after = Pt(8)
    if cover_no == 1:
        section_break(cover_section)


def format_table(table, widths, keep_whole=False):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for col, width in zip(table.columns, widths):
        col.width = Cm(width)
    pr = table._tbl.tblPr
    put(pr, 'w:tblW', **{'w:w': 8504, 'w:type': 'dxa'})
    for node in pr.findall(qn('w:tblInd')):
        pr.remove(node)
    borders = put(pr, 'w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        put(borders, 'w:' + edge, **{'w:val': 'single', 'w:sz': '6', 'w:color': '000000'})
    margins = put(pr, 'w:tblCellMar')
    for edge in ('top', 'bottom'):
        put(margins, 'w:' + edge, **{'w:w': '70', 'w:type': 'dxa'})
    for edge in ('left', 'right'):
        put(margins, 'w:' + edge, **{'w:w': '90', 'w:type': 'dxa'})
    for ri, row in enumerate(table.rows):
        trpr = row._tr.get_or_add_trPr()
        for h in list(trpr.findall(qn('w:trHeight'))):
            trpr.remove(h)
        put(trpr, 'w:cantSplit')
        if ri == 0:
            put(trpr, 'w:tblHeader')
        for cell, width in zip(row.cells, widths):
            cell.width = Cm(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for shading in list(cell._tc.get_or_add_tcPr().findall(qn('w:shd'))):
                shading.getparent().remove(shading)
            for para in cell.paragraphs:
                para.style = doc.styles['ReportBody']
                pf = para.paragraph_format
                pf.first_line_indent = Cm(0)
                pf.space_before = Pt(0)
                pf.space_after = Pt(0)
                pf.line_spacing = 1.3
                pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
                pf.keep_with_next = keep_whole and ri < len(table.rows) - 1
                pf.keep_together = True
                for run in para.runs:
                    set_font(run, 12, ri == 0, False)


p('LỊCH LÀM VIỆC NHÓM THEO TUẦN', 'FrontTitle').paragraph_format.page_break_before = False
schedule = doc.add_table(rows=11, cols=5)
for c, value in zip(schedule.rows[0].cells, ['Tuần', 'Nội dung công việc', 'Thành viên', 'Minh chứng', 'Trạng thái']):
    c.text = value
for week in range(1, 11):
    schedule.cell(week, 0).text = str(week)
format_table(schedule, [1.05, 4.65, 3.05, 3.5, 2.75], True)
for row in schedule.rows[1:]:
    row.height = Cm(1.1)

p('BẢNG PHÂN CÔNG CÔNG VIỆC', 'FrontTitle')
assignment = doc.add_table(rows=4, cols=5)
for c, value in zip(assignment.rows[0].cells, ['STT', 'MSSV', 'Họ và tên', 'Nội dung công việc', 'Mức đóng góp (%)']):
    c.text = value
for i, (mssv, name) in enumerate(MEMBERS, 1):
    assignment.cell(i, 0).text = str(i)
    assignment.cell(i, 1).text = mssv
    assignment.cell(i, 2).text = name
    assignment.cell(i, 3).text = 'Chương 2: Cơ sở lý thuyết và công nghệ sử dụng' if i == 2 else 'Chương 1: Tổng quan về bài toán' if i == 3 else ''
format_table(assignment, [.85, 2.4, 3.4, 5.85, 2.5], True)
for row in assignment.rows[1:]:
    row.height = Cm(1.3)
p('LỜI CẢM ƠN', 'FrontTitle')
p()
p('MỤC LỤC', 'FrontTitleNoTOC')
field(p(style='toc 1'), 'TOC \\h \\z \\t "FrontTitle,1,Heading 11,1,ChapterTitle,1,Heading 21,2,IntroSectionTitle,2,ConclusionSectionTitle,2,Heading 31,3"', 'Mục lục')

p('DANH MỤC CÁC KÝ HIỆU, CHỮ VIẾT TẮT VÀ THUẬT NGỮ', 'FrontTitle')
terms = [
    ('kW', 'Kilowatt', 'Đơn vị công suất; 1 kW = 1.000 W'),
    ('Wh, kWh', 'Watt-hour, kilowatt-hour', 'Đơn vị điện năng; 1 kWh = 1.000 Wh'),
    ('AMI', 'Advanced Metering Infrastructure', 'Hạ tầng đo đếm tiên tiến'),
    ('UCI', 'University of California, Irvine', 'Đơn vị duy trì kho dữ liệu UCI Machine Learning Repository'),
    ('MAE', 'Mean Absolute Error', 'Sai số tuyệt đối trung bình'),
    ('RMSE', 'Root Mean Squared Error', 'Căn bậc hai sai số bình phương trung bình'),
    ('MAPE', 'Mean Absolute Percentage Error', 'Sai số phần trăm tuyệt đối trung bình'),
    ('ARIMA', 'Autoregressive Integrated Moving Average', 'Mô hình tự hồi quy tích hợp trung bình trượt'),
    ('SARIMA', 'Seasonal ARIMA', 'Mô hình ARIMA có thành phần mùa vụ'),
    ('LSTM', 'Long Short-Term Memory', 'Mạng nơ-ron hồi tiếp với bộ nhớ dài–ngắn hạn'),
    ('GRU', 'Gated Recurrent Unit', 'Đơn vị hồi tiếp có cổng'),
    ('HDFS', 'Hadoop Distributed File System', 'Hệ thống tệp phân tán Hadoop'),
    ('Forecasting horizon', 'Forecasting horizon', 'Khoảng từ thời điểm phát hành đến thời điểm cần dự báo'),
    ('Event Time', 'Event Time', 'Thời gian sự kiện theo dấu thời gian dữ liệu'),
    ('Watermark', 'Watermark', 'Dấu mốc tiến độ thời gian sự kiện'),
    ('Window', 'Window', 'Cửa sổ giới hạn phạm vi dữ liệu để tính toán'),
    ('State', 'State', 'Trạng thái được duy trì qua các bản ghi'),
    ('Checkpoint', 'Checkpoint', 'Ảnh chụp trạng thái làm cơ sở khôi phục sau sự cố'),
]
glossary = doc.add_table(rows=1 + len(terms), cols=3)
for c, value in zip(glossary.rows[0].cells, ['Ký hiệu / thuật ngữ', 'Tiếng Anh', 'Nghĩa tiếng Việt']):
    c.text = value
for row, values in zip(glossary.rows[1:], terms):
    for cell, value in zip(row.cells, values):
        cell.text = value
format_table(glossary, [3.0, 5.4, 6.6])
p('DANH MỤC HÌNH', 'FrontTitle')
field(p(style='toc 1'), 'TOC \\h \\z \\t "FigureCaption,1"')
p('DANH MỤC BẢNG', 'FrontTitle')
field(p(style='toc 1'), 'TOC \\h \\z \\t "TableCaption,1"')
section_break(front_section)
p('MỞ ĐẦU', 'Heading 11').paragraph_format.page_break_before = False
for el in intro_prototypes:
    append(el)
    p()
append(deepcopy(chapter_headers[0]))


def h2(value):
    el = deepcopy(h2_prototype)
    para = Paragraph(el, doc._body)
    para.clear()
    set_font(para.add_run(value))
    append(el)
    return para


def caption(label, number, value):
    para = p(style='FigureCaption' if label == 'Hình' else 'TableCaption')
    set_font(para.add_run(label + ' '), 12)
    field(para, 'STYLEREF "ChapterTitle" \\n', '1')
    set_font(para.add_run('.'), 12)
    field(para, ('SEQ Figure' if label == 'Hình' else 'SEQ Table') + ' \\s 1', str(number))
    set_font(para.add_run('. ' + value), 12)
    return para


def mr(value):
    r = OxmlElement('m:r')
    t = OxmlElement('m:t')
    t.text = value
    r.append(t)
    return r


def sub(base, index):
    node = OxmlElement('m:sSub')
    e = OxmlElement('m:e')
    e.append(base)
    s = OxmlElement('m:sub')
    s.append(mr(index))
    node.extend([e, s])
    return node


def acc(value, mark):
    node = OxmlElement('m:acc')
    pr = OxmlElement('m:accPr')
    ch = OxmlElement('m:chr')
    ch.set(qn('m:val'), mark)
    pr.append(ch)
    e = OxmlElement('m:e')
    e.append(mr(value))
    node.extend([pr, e])
    return node


def fraction(numerator, denominator):
    node = OxmlElement('m:f')
    a = OxmlElement('m:num')
    a.append(mr(numerator))
    b = OxmlElement('m:den')
    b.append(mr(denominator))
    node.extend([a, b])
    return node


def sigma(squared=False):
    node = OxmlElement('m:nary')
    pr = OxmlElement('m:naryPr')
    ch = OxmlElement('m:chr')
    ch.set(qn('m:val'), '∑')
    loc = OxmlElement('m:limLoc')
    loc.set(qn('m:val'), 'undOvr')
    pr.extend([ch, loc])
    lo = OxmlElement('m:sub')
    lo.append(mr('i=1'))
    hi = OxmlElement('m:sup')
    hi.append(mr('n'))
    e = OxmlElement('m:e')
    delta = [mr('(' if squared else '|'), sub(mr('y'), 'i'), mr('−'), sub(acc('y', '̂'), 'i'), mr(')' if squared else '|')]
    if squared:
        power = OxmlElement('m:sSup')
        base = OxmlElement('m:e')
        base.extend(delta)
        sup = OxmlElement('m:sup')
        sup.append(mr('2'))
        power.extend([base, sup])
        e.append(power)
    else:
        e.extend(delta)
    node.extend([pr, lo, hi, e])
    return node


def equation(kind):
    para = p()
    para.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.first_line_indent = Cm(0)
    para.paragraph_format.keep_with_next = True
    omathp = OxmlElement('m:oMathPara')
    omath = OxmlElement('m:oMath')
    if kind == 'energy':
        omath.extend([sub(mr('E'), 'i'), mr('='), sub(acc('P', '̅'), 'i'), mr('×'), sub(mr('Δt'), 'i')])
    elif kind == 'mae':
        omath.extend([mr('MAE='), fraction('1', 'n'), mr('×'), sigma()])
    else:
        radical = OxmlElement('m:rad')
        pr = OxmlElement('m:radPr')
        hide = OxmlElement('m:degHide')
        hide.set(qn('m:val'), '1')
        pr.append(hide)
        degree = OxmlElement('m:deg')
        e = OxmlElement('m:e')
        e.extend([fraction('1', 'n'), mr('×'), sigma(True)])
        radical.extend([pr, degree, e])
        omath.extend([mr('RMSE='), radical])
    omathp.append(omath)
    para._p.append(omathp)


added_13 = False
citation_map = {1: [1], 7: [3], 16: [4], 19: [5], 20: [6], 25: [5], 26: [7], 36: [4], 41: [5], 43: [5], 44: [5], 46: [8], 47: [9], 52: [10], 53: [11], 54: [12], 55: [13], 58: [14], 61: [4], 63: [4], 66: [16]}
changes = []
for block_index, el in enumerate(chapter.element.body):
    if block_index == 0:
        continue
    if block_index >= 68:
        break
    if el.tag == qn('w:tbl'):
        tbl = Table(append(deepcopy(el)), doc._body)
        if block_index == 60:
            tbl.cell(1, 2).text = 'HDFS là một hệ thống tệp phân tán [15]'
            tbl.cell(2, 2).text = 'Apache Flink hỗ trợ xử lý dữ liệu hữu hạn và luồng, quản lý trạng thái [16]'
        format_table(tbl, [4.0, 5.5, 5.5], block_index in (12, 30, 60))
        continue
    if el.tag != qn('w:p'):
        continue
    if block_index == 17 and not added_13:
        h2(chapter1_titles[2])
        p('Công tơ thông minh là thành phần đo đếm có khả năng ghi nhận và trao đổi dữ liệu sử dụng điện. Hạ tầng đo đếm tiên tiến (Advanced Metering Infrastructure – AMI) bao gồm công tơ, hệ thống truyền thông và hệ thống quản lý dữ liệu; các thành phần phối hợp để thu thập thông tin đo và hỗ trợ trao đổi hai chiều giữa khách hàng với đơn vị điện lực [2].')
        p('Dữ liệu đo theo những khoảng thời gian xác định cho phép quan sát hồ sơ sử dụng điện chi tiết hơn chỉ số tích lũy theo kỳ. Mỗi bản ghi cần được diễn giải cùng dấu thời gian, đại lượng và đơn vị đo. Độ chi tiết của bản ghi không tự chứng minh dữ liệu đầy đủ hoặc hệ thống đang vận hành theo thời gian thực.')
        p('Bộ UCI được sử dụng trong đề tài chứa các phép đo từng phút của một hộ gia đình và ba nhóm đo phụ. Đây là dữ liệu lịch sử công khai, không phải bằng chứng về một hệ thống AMI hoặc công tơ thời gian thực do nhóm triển khai. Các trường công suất, điện áp, cường độ dòng điện và điện năng đo phụ cung cấp thông tin để khảo sát mức sử dụng điện và xây dựng chuỗi điện năng theo giờ [4].')
        added_13 = True
    if el.xpath('.//w:drawing'):
        para = p()
        para.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.first_line_indent = Cm(0)
        para.paragraph_format.keep_with_next = True
        for drawing in el.xpath('.//w:drawing'):
            copied = deepcopy(drawing)
            remap_images(copied, chapter)
            extent = copied.xpath('.//wp:extent')[0]
            ratio = int(extent.get('cy')) / int(extent.get('cx'))
            width = int(Cm(15))
            extent.set('cx', str(width))
            extent.set('cy', str(int(width * ratio)))
            for aext in copied.xpath('.//a:xfrm/a:ext'):
                aext.set('cx', str(width))
                aext.set('cy', str(int(width * ratio)))
            run = OxmlElement('w:r')
            run.append(copied)
            para._p.append(run)
        caption('Hình', 1, 'Định hướng tiếp cận của đề tài')
        continue
    lines = [line.strip() for line in text(el).split('\n') if line.strip()]
    for line_no, line in enumerate(lines):
        match = re.match(r'^1\.(\d+)\.\s+(.*)$', line)
        if match:
            h2(match.group(2))
            continue
        cap = re.match(r'^Bảng 1\.(\d+)\.\s+(.*)$', line)
        if cap:
            caption('Bảng', int(cap.group(1)), cap.group(2))
            continue
        if block_index in (9, 39, 40):
            equation({9: 'energy', 39: 'mae', 40: 'rmse'}[block_index])
            continue
        line = line.replace('[6]', '[5]').replace(' .', '.').strip()
        if block_index == 13:
            line = 'Nguồn: tổng hợp từ định nghĩa đại lượng và quan hệ giữa công suất với điện năng [3].'
        if block_index == 54:
            line = line.replace('mạng có cổng chỉ vượt trội trong một trường hợp hộ gia đình được xét.', 'ở cấp hộ gia đình, một số cấu hình mạng có cổng cho kết quả nhỉnh hơn nhưng chưa đủ bằng chứng để khẳng định ưu thế chung.')
        if block_index == 66:
            line = line.replace('Phương pháp dự báo và công nghệ lưu trữ hoặc xử lý dữ liệu sẽ được xác định cụ thể ở phần triển khai, bảo đảm phù hợp với dữ liệu và yêu cầu sử dụng công nghệ Big Data của môn học.', 'Apache Flink được lựa chọn cho lớp tiếp nhận và xử lý dữ liệu; bước huấn luyện và đánh giá mô hình dự báo là một thành phần riêng. Phương pháp dự báo cụ thể sẽ được lựa chọn qua thực nghiệm trên biến mục tiêu điện năng của giờ kế tiếp.')
        if block_index in citation_map and line_no == len(lines) - 1:
            missing = [n for n in citation_map[block_index] if f'[{n}]' not in line]
            if missing:
                line = line.rstrip('. ') + ' ' + ', '.join(f'[{n}]' for n in missing) + '.'
        para = p(line)
        if block_index == 13:
            para.paragraph_format.first_line_indent = Cm(0)
            para.paragraph_format.keep_together = True
            for r in para.runs:
                set_font(r, 12)
        changes.append({'source_block': block_index, 'text': line})

replacements = {
    'Giới thiệu bộ dữ liệu London Smart Meter': 'Giới thiệu bộ dữ liệu UCI Individual Household Electric Power Consumption',
    'Phân tích mức tiêu thụ theo hộ sử dụng': 'Phân tích các nhóm đo phụ trong một hộ sử dụng điện',
    'Xác định bài toán dự đoán': 'Xác định bài toán dự đoán điện năng của hộ trong giờ tiếp theo',
    'Xác định biến mục tiêu': 'Xác định biến mục tiêu điện năng theo giờ (kWh)',
    'Xử lý dữ liệu theo cửa sổ thời gian': 'Xử lý dữ liệu theo cửa sổ thời gian một giờ',
    'Tích hợp mô hình dự đoán': 'Tích hợp mô hình đã huấn luyện để thực hiện dự đoán',
}
for el in tail:
    if el.tag == qn('w:p'):
        value = text(el)
        if value in replacements:
            para = Paragraph(el, doc._body)
            para.clear()
            set_font(para.add_run(replacements[value]))
        for node in el.xpath('.//w:bookmarkStart|.//w:bookmarkEnd'):
            node.getparent().remove(node)
    append(el)


refs = [
    ('U.S. Energy Information Administration', '2011', 'Electricity demand changes in predictable patterns', '', 'https://www.eia.gov/todayinenergy/detail.php?id=4190'),
    ('U.S. Department of Energy', '2016', 'Advanced Metering Infrastructure and Customer Systems: Results from the Smart Grid Investment Grant Program', '', 'https://www.energy.gov/sites/prod/files/2016/12/f34/AMI%20Summary%20Report_09-26-16.pdf'),
    ('U.S. Energy Information Administration', 'n.d.', 'Measuring electricity', 'Energy Explained', 'https://www.eia.gov/energyexplained/electricity/measuring-electricity.php'),
    ('G. Hebrail và A. Berard', '2006', 'Individual Household Electric Power Consumption', 'Bộ dữ liệu, UCI Machine Learning Repository. DOI: 10.24432/C58K54', 'https://archive.ics.uci.edu/dataset/235/individual+household+electric+power+consumption'),
    ('R. J. Hyndman và G. Athanasopoulos', '2021', 'Forecasting: Principles and Practice', 'Ấn bản thứ 3, OTexts', 'https://otexts.com/fpp3/'),
    ('U.S. Energy Information Administration', '2020', 'Hourly electricity consumption varies throughout the day and across seasons', '', 'https://www.eia.gov/todayinenergy/detail.php?id=42915'),
    ('National Institute of Standards and Technology', '2014', 'Guidelines for Smart Grid Cybersecurity', 'NISTIR 7628, Revision 1. DOI: 10.6028/NIST.IR.7628r1', 'https://csrc.nist.gov/pubs/ir/7628/r1/final'),
    ('L. Breiman', '2001', 'Random Forests', 'Machine Learning, 45, tr. 5–32. DOI: 10.1023/A:1010933404324', 'https://www.stat.berkeley.edu/~breiman/randomforest2001.pdf'),
    ('S. Hochreiter và J. Schmidhuber', '1997', 'Long Short-Term Memory', 'Neural Computation, 9(8), tr. 1735–1780. DOI: 10.1162/neco.1997.9.8.1735', 'https://www.bioinf.jku.at/publications/older/2604.pdf'),
    ('S. Haben, C. Singleton và P. Grindrod', '2016', 'Analysis and Clustering of Residential Customers Energy Behavioral Demand Using Smart Meter Data', 'IEEE Transactions on Smart Grid, 7(1), tr. 136–144. DOI: 10.1109/TSG.2015.2409786', 'https://doi.org/10.1109/TSG.2015.2409786'),
    ('P. Lusis, K. R. Khalilpour, L. Andrew và A. Liebman', '2017', 'Short-term residential load forecasting: Impact of calendar effects and forecast granularity', 'Applied Energy, 205, tr. 654–669. DOI: 10.1016/j.apenergy.2017.07.114', 'https://doi.org/10.1016/j.apenergy.2017.07.114'),
    ('A. Gasparin, S. Lukovic và C. Alippi', '2022', 'Deep learning for time series forecasting: The electric load case', 'CAAI Transactions on Intelligence Technology, 7(1), tr. 1–25. DOI: 10.1049/cit2.12060', 'https://ietresearch.onlinelibrary.wiley.com/doi/full/10.1049/cit2.12060'),
    ('T. Hong, P. Pinson, S. Fan, H. Zareipour, A. Troccoli và R. J. Hyndman', '2016', 'Probabilistic energy forecasting: Global Energy Forecasting Competition 2014 and beyond', 'International Journal of Forecasting, 32(3), tr. 896–913. DOI: 10.1016/j.ijforecast.2016.02.001', 'https://doi.org/10.1016/j.ijforecast.2016.02.001'),
    ('W. L. Chang và N. Grady', '2019', 'NIST Big Data Interoperability Framework: Volume 1, Definitions', 'NIST SP 1500-1r2. DOI: 10.6028/NIST.SP.1500-1r2', 'https://doi.org/10.6028/NIST.SP.1500-1r2'),
    ('Apache Software Foundation', 'n.d.', 'HDFS Architecture', 'Tài liệu Apache Hadoop', 'https://hadoop.apache.org/docs/stable/hadoop-project-dist/hadoop-hdfs/HdfsDesign.html'),
    ('Apache Software Foundation', 'n.d.', 'What is Apache Flink? — Architecture', 'Tài liệu Apache Flink', 'https://flink.apache.org/what-is-flink/flink-architecture/'),
]
for i, (author, year, title, detail, url) in enumerate(refs, 1):
    para = p()
    para.paragraph_format.first_line_indent = Cm(-.75)
    para.paragraph_format.left_indent = Cm(.75)
    para.paragraph_format.keep_together = True
    set_font(para.add_run(f'[{i}] {author} ({year}). '))
    set_font(para.add_run(title), italic=True)
    set_font(para.add_run('. ' + (detail + '. ' if detail else '') + 'Truy cập ngày 08/10/2026. '))
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), doc.part.relate_to(url, RT.HYPERLINK, is_external=True))
    run = OxmlElement('w:r')
    props = OxmlElement('w:rPr')
    fonts = OxmlElement('w:rFonts')
    fonts.set(qn('w:ascii'), 'Times New Roman')
    fonts.set(qn('w:hAnsi'), 'Times New Roman')
    props.append(fonts)
    put(props, 'w:color', **{'w:val': '000000'})
    put(props, 'w:u', **{'w:val': 'single'})
    run.append(props)
    n = OxmlElement('w:t')
    n.text = url
    run.append(n)
    link.append(run)
    para._p.append(link)

for section in doc.sections:
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(3.5)
    section.right_margin = Cm(2.5)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.header_distance = Cm(1.27)
    section.footer_distance = Cm(1.27)
for para in doc.paragraphs:
    for node in para._p.xpath('./w:bookmarkStart|./w:bookmarkEnd'):
        node.getparent().remove(node)
settings = doc.settings.element
put(settings, 'w:updateFields', **{'w:val': 'true'})
put(settings, 'w:defaultTabStop', **{'w:val': '720'})
doc.core_properties.title = TITLE.replace('\n', ' ')
doc.core_properties.subject = 'Báo cáo Big Data — UCI; Chương 1 và khung Chương 2–5'
doc.core_properties.author = ''
doc.core_properties.last_modified_by = ''
doc.save(OUTPUT)
if any(sha(Path(path)) != digest for path, digest in source_hashes.items()):
    raise RuntimeError('Nguồn đã thay đổi')
QA.mkdir(parents=True, exist_ok=True)
manifest = {'output': str(OUTPUT), 'output_sha256': sha(OUTPUT), 'source_hashes': source_hashes, 'refs': refs, 'changes': changes, 'source_image_sha256': [hashlib.sha256(r.target_part.blob).hexdigest() for r in chapter.part.rels.values() if r.reltype == RT.IMAGE], 'added_1_3': added_13, 'tables': len(doc.tables), 'maths': len(doc.element.xpath('.//m:oMath')), 'inline_images': len(doc.inline_shapes)}
(QA / 'build-manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps({k: manifest[k] for k in ('output', 'output_sha256', 'added_1_3', 'tables', 'maths', 'inline_images')}, ensure_ascii=False))
