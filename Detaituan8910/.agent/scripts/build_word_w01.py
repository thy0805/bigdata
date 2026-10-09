import io
import json
import re
from copy import deepcopy
from hashlib import sha256
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from docx.shared import Cm
from lxml import etree as E
from PIL import Image

from word_w01_content import ACK, INTRO, CH1_REPLACEMENTS, CH2_REPLACE, CH2_EXTRA, CH2_14, CH3, CH4, CH5, CONCLUSION, TABLES, FIGURES, NEW_REFERENCES

ROOT = Path(__file__).resolve().parents[2]
QA = ROOT / '.agent/qa/word-w01-20261010'
SOURCE = ROOT / 'BaoCao_PhanTich_DuDoan_DienNang_UCI_QA_v3_20261008.docx'
HAU = Path('C:/Users/thy/Downloads/HAULYVUNHAN_Chuong2_CoSoLyThuyet_HoanChinh.docx')
OUTPUT = ROOT / 'BaoCao_PhanTich_DuDoan_DienNang_UCI_W01_v1_20261010.docx'
NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main', 'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships', 'm': 'http://schemas.openxmlformats.org/officeDocument/2006/math', 'wp': 'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing'}
REL = 'http://schemas.openxmlformats.org/package/2006/relationships'

def w(name):
    return '{' + NS['w'] + '}' + name

def el(tag, **attrs):
    n = E.Element(w(tag))
    for k, v in attrs.items():
        n.set(w(k), str(v))
    return n

def run(text, bold=False, italic=False, size=None):
    n = el('r')
    if bold or italic or size:
        pr = el('rPr')
        if bold:
            pr.append(el('b'))
        if italic:
            pr.append(el('i'))
        if size:
            pr.append(el('sz', val=size))
        n.append(pr)
    t = el('t')
    t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
    t.text = text
    n.append(t)
    return n

def field(code, result):
    f = el('fldSimple', instr=' ' + code + ' ')
    f.append(run(result))
    return f

def paragraph(text='', style='ReportBody', keep=False):
    p = el('p')
    pr = el('pPr')
    pr.append(el('pStyle', val=style))
    if keep:
        pr.append(el('keepNext'))
    p.append(pr)
    cursor = 0
    for match in re.finditer(r'\{([TF]\d+)\}', text):
        p.append(run(text[cursor:match.start()]))
        key = match.group(1)
        registry = TABLES if key.startswith('T') else FIGURES
        ch, number = registry[key][:2]
        label = 'Bảng' if key.startswith('T') else 'Hình'
        p.append(field(f'REF W01_{key} \\h', f'{label} {ch}.{number}'))
        cursor = match.end()
    if cursor < len(text):
        p.append(run(text[cursor:]))
    return p

def replace_text(p, text):
    pr = deepcopy(p.find(w('pPr')))
    starts = [deepcopy(x) for x in p.findall(w('bookmarkStart'))]
    ends = [deepcopy(x) for x in p.findall(w('bookmarkEnd'))]
    p[:] = []
    if pr is not None:
        p.append(pr)
    p.extend(starts)
    p.append(run(text))
    p.extend(ends)

def ptext(p):
    return ''.join(p.xpath('.//w:t/text()', namespaces=NS))

def style(p):
    v = p.xpath('./w:pPr/w:pStyle/@w:val', namespaces=NS)
    return v[0] if v else ''

def caption(key, label, sequence, chapter, number, title):
    global bookmark_id
    p = paragraph('', 'TableCaption' if sequence == 'Table' else 'FigureCaption', sequence == 'Table')
    p.append(el('bookmarkStart', id=bookmark_id, name='W01_' + key))
    p.append(run(label + ' '))
    p.append(field('STYLEREF "ChapterTitle" \\n \\t', str(chapter)))
    p.append(run('.'))
    p.append(field(f'SEQ {sequence} \\s 1', str(number)))
    p.append(el('bookmarkEnd', id=bookmark_id))
    bookmark_id += 1
    p.append(run('. ' + title))
    return p

def source_note(text):
    p = paragraph(text)
    pr = p.find(w('pPr'))
    pr.append(el('jc', val='left'))
    pr.append(el('ind', left=0, firstLine=0))
    for r in p.findall(w('r')):
        props = el('rPr')
        props.append(el('i'))
        props.append(el('sz', val=22))
        r.insert(0, props)
    return p

def table_blocks(key):
    ch, number, title, headers, widths, rows, source = TABLES[key]
    t = el('tbl')
    pr = el('tblPr')
    pr.append(el('tblW', type='dxa', w=8504))
    pr.append(el('tblLayout', type='fixed'))
    borders = el('tblBorders')
    for name in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        borders.append(el(name, val='single', sz=6, color='000000'))
    pr.append(borders)
    margins = el('tblCellMar')
    for name in ('top', 'bottom'):
        margins.append(el(name, w=60, type='dxa'))
    for name in ('left', 'right'):
        margins.append(el(name, w=90, type='dxa'))
    pr.append(margins)
    t.append(pr)
    grid = el('tblGrid')
    twips = [round(x / 2.54 * 1440) for x in widths]
    for width in twips:
        grid.append(el('gridCol', w=width))
    t.append(grid)
    for idx, row in enumerate([headers] + rows):
        tr = el('tr')
        trpr = el('trPr')
        trpr.append(el('cantSplit'))
        if idx == 0:
            trpr.append(el('tblHeader'))
        tr.append(trpr)
        for text, width in zip(row, twips):
            tc = el('tc')
            tcpr = el('tcPr')
            tcpr.append(el('tcW', w=width, type='dxa'))
            tcpr.append(el('vAlign', val='center'))
            tc.append(tcpr)
            p = paragraph()
            pp = p.find(w('pPr'))
            pp.append(el('ind', left=0, firstLine=0))
            pp.append(el('jc', val='left'))
            pp.append(el('spacing', before=0, after=0, line=312, lineRule='auto'))
            if idx == 0:
                pp.append(el('keepNext'))
            p.append(run(str(text), bold=idx == 0, size=24))
            tc.append(p)
            tr.append(tc)
        t.append(tr)
    return [caption(key, 'Bảng', 'Table', ch, number, title), t, source_note(source)]

def figure_blocks(key):
    global relationship_id, drawing_id
    ch, number, title, section, height, image_id, intended = FIGURES[key]
    d = Document()
    stream = io.BytesIO()
    Image.new('RGB', (1450, round(height * 100)), (0, 0, 0)).save(stream, format='PNG')
    stream.seek(0)
    p = d.add_paragraph()
    p.add_run().add_picture(stream, width=Cm(14.5), height=Cm(height))
    x = E.fromstring(E.tostring(p._p))
    x.insert(0, el('pPr'))
    pp = x.find(w('pPr'))
    pp.append(el('pStyle', val='ReportBody'))
    pp.append(el('jc', val='center'))
    pp.append(el('keepNext'))
    pp.append(el('ind', left=0, firstLine=0))
    rid = f'rIdW01Image{relationship_id}'
    relationship_id += 1
    for blip in x.xpath('.//*[@r:embed]', namespaces=NS):
        blip.set('{' + NS['r'] + '}embed', rid)
    for dp in x.xpath('.//wp:docPr', namespaces=NS):
        dp.set('id', str(drawing_id))
        dp.set('name', f'W01_{image_id}_{key}')
        dp.set('descr', 'Khung đen chờ ảnh thật W02: ' + intended)
        drawing_id += 1
    media = 'media/w01_' + key + '.png'
    packages['word/' + media] = stream.getvalue()
    rel = E.SubElement(relationships, '{' + REL + '}Relationship')
    rel.set('Id', rid)
    rel.set('Type', NS['r'] + '/image')
    rel.set('Target', media)
    registry.append({'key': key, 'image_id': image_id, 'caption': f'Hình {ch}.{number}. {title}', 'section': section, 'width_cm': 14.5, 'height_cm': height, 'status': 'PLACEHOLDER_W02_PENDING', 'intended_source': intended, 'bookmark': 'W01_' + key, 'media': media})
    return [x, caption(key, 'Hình', 'Figure', ch, number, title)]

def blocks(text):
    result = []
    for part in text.split('\n\n'):
        if part.startswith('@TABLE:'):
            result.extend(table_blocks(part.split(':', 1)[1]))
        elif part.startswith('@FIG:'):
            result.extend(figure_blocks(part.split(':', 1)[1]))
        else:
            result.append(paragraph(part))
    return result

def insert_after(anchor, new):
    index = body.index(anchor) + 1
    while index < len(body):
        candidate = body[index]
        if candidate.tag == w('p') and style(candidate) == 'ReportBody' and not ptext(candidate) and not candidate.xpath('.//w:br|.//w:sectPr|.//w:drawing|.//w:instrText|.//w:fldSimple', namespaces=NS):
            body.remove(candidate)
        else:
            break
    for offset, item in enumerate(new):
        body.insert(index + offset, item)

def reference(number, description, url):
    global relationship_id
    p = paragraph(f'[{number}] {description} Truy cập ngày 10/10/2026. ', 'ReportBody')
    pp = p.find(w('pPr'))
    pp.append(el('keepLines'))
    pp.append(el('ind', left=425, hanging=425))
    pp.append(el('jc', val='left'))
    rid = f'rIdW01Link{relationship_id}'
    relationship_id += 1
    link = E.SubElement(relationships, '{' + REL + '}Relationship')
    link.set('Id', rid)
    link.set('Type', NS['r'] + '/hyperlink')
    link.set('Target', url)
    link.set('TargetMode', 'External')
    h = el('hyperlink')
    h.set('{' + NS['r'] + '}id', rid)
    r = run(url)
    rp = el('rPr')
    rp.append(el('rStyle', val='Hyperlink'))
    r.insert(0, rp)
    h.append(r)
    p.append(h)
    return p

if OUTPUT.exists():
    raise SystemExit('Refusing to overwrite existing W01 output')
assert sha256(SOURCE.read_bytes()).hexdigest() == 'efc3acdfec05bc264eb3320f2df363241cab29818830a7159cc88aabdb77bdd1'
assert sha256(HAU.read_bytes()).hexdigest() == '0902902041a6426555b3ce852e188ef3948c61a39e7cb0bf087c5aa85638065d'
with ZipFile(SOURCE) as z:
    packages = {name: z.read(name) for name in z.namelist()}
document = E.fromstring(packages['word/document.xml'])
body = document.find(w('body'))
original = list(body)
relationships = E.fromstring(packages['word/_rels/document.xml.rels'])
bookmark_id = max([int(v) for v in document.xpath('.//w:bookmarkStart/@w:id', namespaces=NS)] + [0]) + 1
relationship_id = 1
drawing_id = max([int(v) for v in document.xpath('.//wp:docPr/@id', namespaces=NS)] + [0]) + 1
registry = []
changelog = []
equations = [deepcopy(x) for x in original if x.xpath('.//m:oMath', namespaces=NS)]
chapter = 0
sections = {}
intro = []
conclusion = []
for node in original:
    s = style(node)
    if s == 'ChapterTitle':
        chapter += 1
        sections[chapter] = []
    elif s == 'Heading21':
        sections[chapter].append(node)
    elif s == 'IntroSectionTitle':
        intro.append(node)
    elif s == 'ConclusionSectionTitle':
        conclusion.append(node)
assert [len(sections[c]) for c in range(1, 6)] == [11, 14, 13, 14, 17]
assert len(intro) == 5 and len(conclusion) == 3
for index, text in CH1_REPLACEMENTS.items():
    node = original[index - 1]
    assert style(node) == 'ReportBody'
    changelog.append({'chapter': 1, 'block': index, 'before': ptext(node), 'after': text, 'reason': 'Đồng bộ dự kiến với thực nghiệm đã kiểm; giữ phần còn lại.'})
    replace_text(node, text)
for index, title in [(6, 'Nhóm dữ liệu theo khóa thời gian trong SQL'), (7, 'Xử lý dấu thời gian trong Flink SQL BATCH'), (8, 'Tổng hợp dữ liệu theo khung giờ bằng SQL')]:
    node = sections[5][index - 1]
    changelog.append({'chapter': 5, 'section': f'5.{index}', 'before': ptext(node), 'after': title, 'reason': 'Mô tả đúng GROUP BY/FLOOR trong SQL BATCH.'})
    replace_text(node, title)
ack_title = next(x for x in original if ptext(x) == 'LỜI CẢM ƠN' and style(x) not in ('TOC1', 'TOC2'))
insert_after(ack_title, blocks(ACK))
for idx, node in enumerate(intro, 1):
    insert_after(node, blocks(INTRO[idx]))
with ZipFile(HAU) as z:
    hb = E.fromstring(z.read('word/document.xml')).find(w('body'))
hparts = {i: [] for i in range(1, 15)}
current = 0
for bid, node in enumerate(hb, 1):
    txt = ptext(node)
    match = re.match(r'^2\.(\d+)\.', txt)
    if match:
        current = int(match.group(1))
        continue
    if not current or node.tag != w('p') or not txt:
        continue
    if current == 14:
        continue
    revised = CH2_REPLACE.get(bid, txt)
    if bid in CH2_REPLACE:
        changelog.append({'chapter': 2, 'section': f'2.{current}', 'source_block': bid, 'before': txt, 'after': revised, 'reason': 'Sửa có chọn lọc theo CHAPTER2_REVIEW và hiện thực đã khóa.'})
    outstyle = 'ReportBullet' if style(node) == 'ListBullet' else ('ReportSubhead' if len(revised) < 115 and not revised.endswith('.') else 'ReportBody')
    hparts[current].append(paragraph(revised, outstyle))
    if current == 8 and bid in (84, 86):
        eq = deepcopy(equations[1 if bid == 84 else 2])
        for bm in eq.xpath('.//w:bookmarkStart|.//w:bookmarkEnd', namespaces=NS):
            bm.getparent().remove(bm)
        hparts[current].append(eq)
for number, extra in CH2_EXTRA.items():
    hparts[number].append(paragraph(extra))
hparts[14] = blocks(CH2_14)
changelog.append({'chapter': 2, 'section': '2.14', 'reason': 'Thay mô tả stack Kafka/Grafana chưa triển khai bằng Flink SQL BATCH/Python/Streamlit/Plotly theo mẫu đã duyệt.'})
for index, node in enumerate(sections[2], 1):
    assert hparts[index], index
    insert_after(node, hparts[index])
for number, content in [(3, CH3), (4, CH4), (5, CH5)]:
    for index, node in enumerate(sections[number], 1):
        insert_after(node, blocks(content[index]))
for index, node in enumerate(conclusion, 1):
    insert_after(node, blocks(CONCLUSION[index]))
for number, (description, url) in enumerate(NEW_REFERENCES, 17):
    body.insert(len(body) - 1, reference(number, description, url))
settings = E.fromstring(packages['word/settings.xml'])
update = settings.find(w('updateFields'))
if update is None:
    update = el('updateFields', val='true')
    settings.append(update)
else:
    update.set(w('val'), 'true')
packages['word/settings.xml'] = E.tostring(settings, encoding='UTF-8', xml_declaration=True, standalone=True)
packages['word/document.xml'] = E.tostring(document, encoding='UTF-8', xml_declaration=True, standalone=True)
packages['word/_rels/document.xml.rels'] = E.tostring(relationships, encoding='UTF-8', xml_declaration=True, standalone=True)
with ZipFile(OUTPUT, 'w', ZIP_DEFLATED) as z:
    for name, content in packages.items():
        z.writestr(name, content)
QA.mkdir(parents=True, exist_ok=True)
(QA / 'image-registry.json').write_text(json.dumps({'status': 'PLACEHOLDER_ONLY', 'omitted': ['IMG01: bảng 9 cột đã giải thích dữ liệu gốc'], 'images': registry}, ensure_ascii=False, indent=2), encoding='utf-8')
(QA / 'content-changelog.json').write_text(json.dumps(changelog, ensure_ascii=False, indent=2), encoding='utf-8')
(QA / 'authored.json').write_text(json.dumps({'status': 'APPLIED_UNVERIFIED', 'output': str(OUTPUT), 'sha256_before_word_update': sha256(OUTPUT.read_bytes()).hexdigest(), 'new_tables': len(TABLES), 'new_black_placeholders': len(registry), 'source_parts_unchanged': [name for name in packages if name not in ('word/document.xml', 'word/settings.xml', 'word/_rels/document.xml.rels') and not name.startswith('word/media/w01_')], 'chapter2_sections_imported': 14, 'replaced_hau_blocks': len(CH2_REPLACE)}, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps({'output': str(OUTPUT), 'status': 'APPLIED_UNVERIFIED', 'tables_added': len(TABLES), 'placeholders': len(registry)}, ensure_ascii=False))
