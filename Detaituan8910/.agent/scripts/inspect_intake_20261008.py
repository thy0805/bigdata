import hashlib
import json
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from lxml import etree as E
from pypdf import PdfReader
import pypdfium2 as pdfium

root = Path(r'D:\Hoctap\bigdata')
project = root / 'Detaituan8910'
qa = project / '.agent/qa/intake-20261008'
qa.mkdir(parents=True, exist_ok=True)
sources = {
    'teacher': project / 'Ke hoach Do an mon hoc Nhap mon Big data.pdf',
    'chapter1': project / 'KhoiTuan Tong Quan Bai Toan Chuong 1.docx',
    'chapter2': project / 'HAULYVUNHAN_Chuong2_CoSoLyThuyet.docx',
    'skeleton': root / 'BaoCao_PhanTich_DuDoan_DienNang_Khung_v2.docx',
    'template': root / 'Mauwword/Mau bao cao_Do an_Khoa luan_2025.docx',
}
ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main', 'm': 'http://schemas.openxmlformats.org/officeDocument/2006/math'}
result = {}
for name, path in sources.items():
    data = path.read_bytes()
    item = {'path': str(path), 'sha256': hashlib.sha256(data).hexdigest(), 'bytes': len(data)}
    lines = []
    if path.suffix.lower() == '.pdf':
        reader = PdfReader(path)
        item['pages'] = len(reader.pages)
        for n, page in enumerate(reader.pages, 1):
            lines += [f'PAGE {n}', page.extract_text(extraction_mode='layout')]
        rendered = pdfium.PdfDocument(str(path))
        folder = qa / f'{name}-render'
        folder.mkdir(exist_ok=True)
        for n in range(len(rendered)):
            rendered[n].render(scale=2).to_pil().save(folder / f'page-{n + 1}.png')
        rendered.close()
    else:
        doc = Document(path)
        item['paragraphs'] = len(doc.paragraphs)
        item['tables'] = len(doc.tables)
        item['inline_shapes'] = len(doc.inline_shapes)
        item['sections'] = [{'width_cm': s.page_width.cm, 'height_cm': s.page_height.cm, 'top_cm': s.top_margin.cm, 'bottom_cm': s.bottom_margin.cm, 'left_cm': s.left_margin.cm, 'right_cm': s.right_margin.cm} for s in doc.sections]
        with ZipFile(path) as z:
            xml = E.fromstring(z.read('word/document.xml'))
            item['math_count'] = len(xml.xpath('//m:oMath', namespaces=ns))
            item['fields'] = xml.xpath('//w:instrText/text()', namespaces=ns)
            item['relationships'] = z.read('word/_rels/document.xml.rels').decode('utf-8')
            item['media'] = [p for p in z.namelist() if p.startswith('word/media/')]
            body = xml.find('w:body', namespaces=ns)
            for index, child in enumerate(body):
                if E.QName(child).localname == 'tbl':
                    lines.append(f'BLOCK {index} TABLE')
                    for row in child.findall('w:tr', namespaces=ns):
                        lines.append(' | '.join(''.join(c.xpath('.//w:t/text() | .//m:t/text()', namespaces=ns)) for c in row.findall('w:tc', namespaces=ns)))
                else:
                    text = ''.join(child.xpath('.//w:t/text() | .//m:t/text()', namespaces=ns))
                    if text:
                        lines.append(f'BLOCK {index}: {text}')
    (qa / f'{name}-text.txt').write_text('\n'.join(lines), encoding='utf-8')
    result[name] = item
(qa / 'sources.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps({k: {p: v for p, v in value.items() if p not in ['relationships', 'fields', 'sections']} for k, value in result.items()}, ensure_ascii=False, indent=2))
