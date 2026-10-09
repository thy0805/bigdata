from pathlib import Path
from docx import Document
from zipfile import ZipFile
from lxml import etree as E
import pypdfium2 as pdfium
import hashlib
import json

root = Path('D:/Hoctap/bigdata/Detaituan8910')
qa = root / '.agent/qa/chapter-outlines'
records = json.loads((qa / 'verification.json').read_text(encoding='utf-8'))
ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main', 'cp': 'http://schemas.openxmlformats.org/package/2006/metadata/core-properties', 'dc': 'http://purl.org/dc/elements/1.1/'}
for record in records:
    number = record['chapter']
    output = Path(record['file'])
    assert hashlib.sha256(output.read_bytes()).hexdigest() == record['sha256']
    doc = Document(output)
    with ZipFile(output) as z:
        xml = E.fromstring(z.read('word/document.xml'))
        assert not xml.xpath('//w:tbl | //w:hyperlink | //w:drawing | //w:commentRangeStart', namespaces=ns)
        assert not any('comments' in name for name in z.namelist())
        core = E.fromstring(z.read('docProps/core.xml'))
        assert not core.xpath('//dc:creator/text() | //cp:lastModifiedBy/text()', namespaces=ns)
    pdf = pdfium.PdfDocument(qa / f'chapter{number}/chapter{number}.pdf')
    assert len(pdf) == 1, (number, len(pdf))
    page = pdf[0]
    page.render(scale=1.7).to_pil().save(qa / f'chapter{number}/page-1.png')
    textpage = page.get_textpage()
    rendered = ' '.join(textpage.get_text_range().split())
    for paragraph in doc.paragraphs:
        assert ' '.join(paragraph.text.split()) in rendered
    record['rendered_pages'] = 1
    record['pdf_contains_all_headings'] = True
    record['source_unchanged_by_word'] = True
    record['no_personal_metadata_or_extra_objects'] = True
(qa / 'verification.json').write_text(json.dumps(records, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps(records, ensure_ascii=False))
