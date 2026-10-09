from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
import hashlib
import json
import re
import unicodedata

root = Path('D:/Hoctap/bigdata/Detaituan8910')
text = (root / 'context.md').read_text(encoding='utf-8')
qa = root / '.agent/qa/chapter-outlines'
qa.mkdir(parents=True, exist_ok=True)
results = []
for number, filename, count in [(1, 'Chuong1_TongQuanBaiToan.docx', 11), (2, 'Chuong2_CoSoLyThuyet.docx', 14)]:
    block = text.split(f'### CHƯƠNG {number}. ', 1)[1].split('\n### ', 1)[0]
    lines = [line.strip() for line in block.splitlines() if line.strip()]
    title = f'CHƯƠNG {number}. ' + lines[0]
    items = lines[1:]
    assert len(items) == count
    assert all(re.match(rf'^{number}\.{i}\. ', value) for i, value in enumerate(items, 1))
    output = root / filename
    if output.exists():
        prior = json.loads((qa / 'verification.json').read_text(encoding='utf-8'))
        expected = next(record['sha256'] for record in prior if record['file'] == str(output))
        assert hashlib.sha256(output.read_bytes()).hexdigest() == expected
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(3.5)
    section.right_margin = Cm(2.5)
    for name, size, bold in [('Normal', 13, False), ('Title', 14, True), ('Heading 1', 13, True)]:
        style = doc.styles[name]
        style.font.name = 'Times New Roman'
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor(0, 0, 0)
        fonts = style.element.get_or_add_rPr().rFonts
        for attr in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
            fonts.set(qn('w:' + attr), 'Times New Roman')
        for attr in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme'):
            fonts.attrib.pop(qn('w:' + attr), None)
        style.paragraph_format.line_spacing = 1.15
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.keep_with_next = False
        style.paragraph_format.keep_together = True
    doc.styles['Title'].paragraph_format.space_after = Pt(16)
    for border in doc.styles.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
    p = doc.add_paragraph(unicodedata.normalize('NFC', title), style='Title')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    for item in items:
        doc.add_paragraph(unicodedata.normalize('NFC', item), style='Heading 1')
    doc.core_properties.author = ''
    doc.core_properties.last_modified_by = ''
    doc.core_properties.title = title
    doc.core_properties.subject = ''
    doc.core_properties.comments = ''
    doc.save(output)
    reopened = Document(output)
    assert [p.text for p in reopened.paragraphs] == [title, *items]
    assert not reopened.tables
    results.append({'file': str(output), 'chapter': number, 'item_count': count, 'only_requested_headings': True, 'font': 'Times New Roman', 'sha256': hashlib.sha256(output.read_bytes()).hexdigest()})
(qa / 'verification.json').write_text(json.dumps(results, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps(results, ensure_ascii=False))
